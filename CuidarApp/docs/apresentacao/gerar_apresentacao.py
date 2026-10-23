"""Monta a apresentação do CuidarApp (PPTX) a partir das capturas em ./telas.

Uso: python3 capturar_telas.py && python3 gerar_apresentacao.py
"""
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

BASE = Path(__file__).parent
TELAS = BASE / "telas"
MOCKUPS = BASE / "mockups"
MOCKUPS.mkdir(exist_ok=True)
OUTPUT = BASE / "CuidarApp_Apresentacao.pptx"

PRIMARY = RGBColor(0x7C, 0x5C, 0xFC)
PRIMARY_DARK = RGBColor(0x5B, 0x3F, 0xD9)
PRIMARY_LIGHT = RGBColor(0xED, 0xE9, 0xFE)
TEXT = RGBColor(0x1F, 0x29, 0x37)
TEXT_2 = RGBColor(0x6B, 0x72, 0x80)
BG = RGBColor(0xFA, 0xFA, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE5, 0xE7, 0xEB)
GREEN = RGBColor(0x10, 0xB9, 0x81)
GREEN_LIGHT = RGBColor(0xD1, 0xFA, 0xE5)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
BLUE_LIGHT = RGBColor(0xE8, 0xF4, 0xFC)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
AMBER_LIGHT = RGBColor(0xFE, 0xF3, 0xC7)
RED = RGBColor(0xDC, 0x26, 0x26)
RED_LIGHT = RGBColor(0xFE, 0xE2, 0xE2)
WATER = RGBColor(0x02, 0x84, 0xC7)
WATER_LIGHT = RGBColor(0xE0, 0xF2, 0xFE)

FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SLIDE_COUNT = [0]


# ==================== MOCKUP DE CELULAR ====================
def mockup(name):
    """Envolve a captura numa moldura de celular com cantos arredondados."""
    out = MOCKUPS / f"{name}.png"
    if out.exists() and out.stat().st_mtime > (TELAS / f"{name}.png").stat().st_mtime:
        return out
    screen = Image.open(TELAS / f"{name}.png").convert("RGBA")
    w, h = screen.size
    bezel = 22
    radius_screen = 70
    radius_frame = radius_screen + bezel
    frame = Image.new("RGBA", (w + bezel * 2, h + bezel * 2), (0, 0, 0, 0))
    draw = ImageDraw.Draw(frame)
    draw.rounded_rectangle([0, 0, frame.width - 1, frame.height - 1], radius=radius_frame, fill=(17, 24, 39, 255))
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, w - 1, h - 1], radius=radius_screen, fill=255)
    frame.paste(screen, (bezel, bezel), mask)
    frame.save(out)
    return out


# ==================== HELPERS ====================
def new_slide(bg=BG):
    slide = prs.slides.add_slide(BLANK)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    SLIDE_COUNT[0] += 1
    return slide


def text(slide, x, y, w, h, content, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    """content pode ser str ou lista de str / (str, dict) para parágrafos."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    paragraphs = content if isinstance(content, list) else [content]
    for i, item in enumerate(paragraphs):
        opts = {}
        if isinstance(item, tuple):
            item, opts = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.line_spacing = line_spacing
        p.space_after = Pt(opts.get("space_after", 0))
        run = p.add_run()
        run.text = item
        f = run.font
        f.name = FONT
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.color.rgb = opts.get("color", color)
    return box


def rect(slide, x, y, w, h, fill=WHITE, radius=0.12, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = min(0.5, radius / min(w, h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def circle(slide, x, y, d, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def phone(slide, name, x, y, height):
    path = mockup(name)
    img = Image.open(path)
    width = height * img.width / img.height
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(height))
    return width


def header(slide, kicker, title, subtitle=None):
    text(slide, 0.7, 0.55, 8, 0.3, kicker.upper(), size=12, color=PRIMARY, bold=True)
    text(slide, 0.7, 0.85, 11.5, 0.8, title, size=32, bold=True)
    if subtitle:
        text(slide, 0.7, 1.6, 11.5, 0.6, subtitle, size=15, color=TEXT_2)


def footer(slide):
    text(slide, 0.7, 7.0, 6, 0.3, "CuidarApp • Cuidado integral do idoso", size=10, color=TEXT_2)
    text(slide, 11.6, 7.0, 1.0, 0.3, str(SLIDE_COUNT[0]), size=10, color=TEXT_2, align=PP_ALIGN.RIGHT)


def feature_list(slide, x, y, w, items, gap=0.12, title_size=16, body_size=13):
    """items: (titulo, descricao, cor_fundo, cor_destaque)."""
    cy = y
    chars_per_line = int((w - 0.5) * 11.5)
    for title, body, light, strong in items:
        lines = -(-len(body) // chars_per_line)
        h = 0.62 + lines * 0.24
        rect(slide, x, cy, w, h, fill=WHITE, radius=0.12, line=BORDER)
        rect(slide, x, cy, 0.09, h, fill=strong, radius=0.04)
        text(slide, x + 0.28, cy + 0.13, w - 0.45, 0.3, title, size=title_size, bold=True)
        text(slide, x + 0.28, cy + 0.45, w - 0.45, h - 0.5, body, size=body_size, color=TEXT_2)
        cy += h + gap
    return cy


def screen_slide(kicker, title, subtitle, screens, items):
    slide = new_slide()
    header(slide, kicker, title, subtitle)
    phone_h = 4.75
    x = 0.7
    for name in screens:
        w = phone(slide, name, x, 2.1, phone_h)
        x += w + 0.3
    feature_list(slide, x + 0.25, 2.1, 13.333 - x - 0.95, items)
    footer(slide)
    return slide


# ==================== 1. CAPA ====================
s = new_slide(bg=PRIMARY)
deco = circle(s, 8.6, -1.6, 6.5, PRIMARY_DARK)
circle(s, 10.9, 4.6, 4.2, RGBColor(0x8F, 0x74, 0xFD))
rect(s, 0.9, 1.15, 0.95, 0.95, fill=WHITE, radius=0.25)
heart = s.shapes.add_shape(MSO_SHAPE.HEART, Inches(1.11), Inches(1.4), Inches(0.53), Inches(0.48))
heart.fill.solid()
heart.fill.fore_color.rgb = PRIMARY
heart.line.fill.background()
heart.shadow.inherit = False
text(s, 0.9, 2.45, 7.5, 1.0, "CuidarApp", size=60, color=WHITE, bold=True)
text(s, 0.9, 3.55, 7.0, 1.2, "Cuidado integral do idoso: remédios, alimentação, hidratação, exercícios e segurança em um só aplicativo.", size=20, color=WHITE)
text(s, 0.9, 5.0, 7.0, 0.4, "100% offline  •  Um aparelho, dois perfis  •  Baseado em evidências clínicas", size=14, color=PRIMARY_LIGHT, bold=True)
text(s, 0.9, 6.55, 7.0, 0.4, "Apresentação de funcionalidades • Versão 1.3", size=12, color=PRIMARY_LIGHT)
phone(s, "01-inicio", 8.0, 0.75, 6.0)
phone(s, "02-idoso-hoje", 10.6, 1.05, 5.4)

# ==================== 2. O DESAFIO ====================
s = new_slide()
header(s, "O desafio", "O Brasil envelhece rápido e o cuidado é fragmentado",
       "Dados reunidos na pesquisa \"Monitoramento Nutricional e Exercício Físico na Saúde do Idoso\".")
stats = [
    ("15,8%", "da população tem 60 anos ou mais", "32,1 milhões de pessoas (Censo IBGE 2022)", PRIMARY, PRIMARY_LIGHT),
    ("15–18%", "dos idosos na comunidade têm sarcopenia", "Perda de massa e força muscular (EWGSOP2)", BLUE, BLUE_LIGHT),
    ("40,9%", "são insuficientemente ativos", "Sedentarismo eleva o risco de quedas (VIGITEL 2023)", GREEN, GREEN_LIGHT),
    ("22–28%", "estão em risco de desnutrição", "Idosos ambulatoriais (IBRANUTRI)", AMBER, AMBER_LIGHT),
]
cw, gap = 2.82, 0.2
for i, (value, label, source, strong, light) in enumerate(stats):
    x = 0.7 + i * (cw + gap)
    rect(s, x, 2.45, cw, 2.6, fill=WHITE, line=BORDER)
    rect(s, x, 2.45, cw, 0.12, fill=strong, radius=0.05)
    text(s, x + 0.3, 2.85, cw - 0.6, 0.8, value, size=40, color=strong, bold=True)
    text(s, x + 0.3, 3.7, cw - 0.6, 0.8, label, size=15, bold=True)
    text(s, x + 0.3, 4.55, cw - 0.6, 0.45, source, size=11, color=TEXT_2)
rect(s, 0.7, 5.35, 11.9, 1.35, fill=PRIMARY_LIGHT)
text(s, 1.0, 5.55, 11.3, 1.0, [
    ("O problema", {"bold": True, "size": 15, "color": PRIMARY_DARK, "space_after": 4}),
    "Apps existentes cuidam de uma parte só (remédio, caloria ou exercício), exigem internet e têm telas densas: "
    "a taxa de abandono passa de 70% entre usuários com mais de 65 anos.",
], size=14)
footer(s)

# ==================== 3. A SOLUÇÃO ====================
s = new_slide()
header(s, "A solução", "Um bloco de notas de cuidados, simples e interativo",
       "O CuidarApp junta tudo o que importa na rotina do idoso em uma tela clara — e funciona sem internet.")
pillars = [
    ("Um aparelho, dois perfis", "O idoso entra direto. O cuidador acessa com senha de 4 dígitos para cadastrar e acompanhar.", PRIMARY, PRIMARY_LIGHT, "👥"),
    ("A tríade da saúde", "Proteína nas refeições, água ao longo do dia e exercício multicomponente, como recomendam ESPEN e OMS.", GREEN, GREEN_LIGHT, "🥗"),
    ("Segurança e SOS", "Botão de emergência sempre visível, registro de quedas e alerta de mal-estar após exercício.", RED, RED_LIGHT, "🚨"),
    ("Offline e privado", "Sem cadastro, sem nuvem. Os dados de saúde ficam somente no celular (Privacy by Design / LGPD).", WATER, WATER_LIGHT, "🔒"),
]
cw = 2.82
for i, (title, body, strong, light, emoji) in enumerate(pillars):
    x = 0.7 + i * (cw + 0.2)
    rect(s, x, 2.5, cw, 3.9, fill=WHITE, line=BORDER)
    circle(s, x + 0.3, 2.8, 0.9, light)
    text(s, x + 0.3, 2.8, 0.9, 0.9, emoji, size=26, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.3, 3.95, cw - 0.6, 0.7, title, size=18, bold=True, color=strong)
    text(s, x + 0.3, 4.7, cw - 0.6, 1.6, body, size=13, color=TEXT_2)
footer(s)

# ==================== 4..12. TELAS ====================
screen_slide("Acesso", "Tela inicial: quem está usando?",
             "Nome e foto de cada perfil (ícone quando não há foto). Nada de e-mail ou login online.",
             ["01-inicio", "10-senha-cuidador"], [
    ("Idoso entra com um toque", "Vai direto para as tarefas do dia, sem senha para lembrar.", GREEN_LIGHT, GREEN),
    ("Cuidador protegido por senha", "PIN de 4 dígitos (padrão 1234), guardado como hash SHA-256 e alterável nas configurações.", PRIMARY_LIGHT, PRIMARY),
    ("Celular do idoso ou do cuidador", "Se o idoso não tem celular ou não consegue usar sozinho, o cuidador ativa o modo \"celular do cuidador\" e registra por ele.", BLUE_LIGHT, BLUE),
    ("Feito para funcionar offline", "Aviso na tela: os dados ficam somente neste aparelho.", WATER_LIGHT, WATER),
])

screen_slide("Perfil do idoso", "\"Hoje\": tudo o que importa em uma tela",
             "Saudação pelo horário, data por extenso e categorias grandes para navegar.",
             ["02-idoso-hoje", "03-idoso-hoje-tarefas"], [
    ("Anel de progresso do dia", "Mostra quanto da rotina já foi cumprida e quantas tarefas faltam.", PRIMARY_LIGHT, PRIMARY),
    ("Copos de água", "Um toque no + a cada copo. A meta diária é definida pelo cuidador.", WATER_LIGHT, WATER),
    ("\"Como você está hoje?\"", "Lembrete para o check-in de humor, que informa o cuidador.", AMBER_LIGHT, AMBER),
    ("Tarefas em ordem de horário", "Remédios, exercícios e refeições com botão grande para marcar como feito. Concluídas vão para o fim da lista.", GREEN_LIGHT, GREEN),
])

screen_slide("Perfil do idoso", "Remédios e alimentação",
             "O idoso vê o que tomar e o que comer hoje; o cuidador acompanha a adesão.",
             ["04-idoso-remedios", "07-idoso-alimentacao"], [
    ("Remédios do dia", "Nome, dose, horário e instruções (\"tomar com água\"). Frequência diária, dias úteis ou semanal.", BLUE_LIGHT, BLUE),
    ("Cardápio do dia", "A dieta semanal montada pelo cuidador aparece automaticamente, refeição por refeição.", AMBER_LIGHT, AMBER),
    ("Meta de proteína", "Calculada pelo peso: 1,0 a 1,2 g/kg, ou 1,2 a 1,5 g/kg em idosos frágeis (ESPEN / PROT-AGE).", GREEN_LIGHT, GREEN),
    ("Hidratação", "Dica para beber água mesmo sem sede: a hipodipsia senil aumenta riscos de confusão e infecção.", WATER_LIGHT, WATER),
])

screen_slide("Perfil do idoso", "Exercício multicomponente com segurança",
             "Força, equilíbrio, marcha e flexibilidade — a combinação que reduz quedas em 23% a 34% (Cochrane).",
             ["05-idoso-exercicios", "06-feedback-exercicio"], [
    ("Meta semanal visível", "Dias ativos na semana em relação à meta de 3 a 5 vezes (OMS / Vivifrail).", GREEN_LIGHT, GREEN),
    ("Exercícios com tipo e instruções", "Cada exercício mostra duração, horário, como fazer e a etiqueta do tipo.", PRIMARY_LIGHT, PRIMARY),
    ("\"Como você se sentiu?\"", "Após marcar o exercício, o idoso responde em um toque. Tontura ou dor geram alerta para o cuidador.", RED_LIGHT, RED),
    ("Atalho para o SOS", "Se não estiver bem, o app oferece ligar imediatamente.", RED_LIGHT, RED),
])

screen_slide("Perfil do idoso", "Como estou e emergência",
             "Saúde emocional e resposta rápida em caso de queda ou mal-estar.",
             ["08-idoso-como-estou", "09-emergencia"], [
    ("Humor em 5 carinhas", "Check-in diário simples. Humor baixo avisa o cuidador (sintomas depressivos caem 46,5% com acompanhamento).", AMBER_LIGHT, AMBER),
    ("Sintomas do dia", "Dor, tontura, cansaço, falta de apetite, sono ruim ou sede, marcados com um toque.", PRIMARY_LIGHT, PRIMARY),
    ("\"Tive uma queda\"", "Registra a queda, alerta o cuidador e, se for preciso, abre a tela de emergência.", RED_LIGHT, RED),
    ("SOS sempre à mão", "Botão central na barra inferior: liga para o 192 ou direto para o cuidador.", RED_LIGHT, RED),
])

screen_slide("Perfil do cuidador", "Painel do cuidador",
             "Visão do dia em segundos: alertas, ações rápidas e progresso.",
             ["11-cuidador-painel", "12-cuidador-progresso"], [
    ("Alertas em destaque", "Quedas, mal-estar no exercício e humor baixo aparecem no topo até serem marcados como vistos.", RED_LIGHT, RED),
    ("Ações rápidas", "Cadastrar medicamento, exercício, visita ou editar a dieta em um toque.", PRIMARY_LIGHT, PRIMARY),
    ("Progresso de hoje", "Remédios, exercícios, refeições e água com barras de progresso.", GREEN_LIGHT, GREEN),
    ("Humor e sintomas", "O que o idoso relatou hoje, sem precisar perguntar.", AMBER_LIGHT, AMBER),
])

screen_slide("Perfil do cuidador", "Acompanhamento semanal e metas",
             "Monitoramento digital eleva a adesão terapêutica de 42,5% (papel) para 78,4%.",
             ["13-cuidador-semana", "15-lista-exercicios"], [
    ("Adesão dos últimos 7 dias", "Gráfico por dia, com cores para dias bons, médios e fracos, e o humor de cada dia.", PRIMARY_LIGHT, PRIMARY),
    ("Indicadores-chave", "Adesão média, dias com exercício e média de copos de água por dia.", WATER_LIGHT, WATER),
    ("Meta de proteína", "Com o peso informado, o app mostra a meta diária em gramas.", AMBER_LIGHT, AMBER),
    ("Treino completo", "A lista de exercícios indica quais tipos ainda faltam para um treino multicomponente.", GREEN_LIGHT, GREEN),
])

screen_slide("Perfil do cuidador", "Cadastro rápido e guiado",
             "Formulários curtos, com sugestões prontas e dicas baseadas nas diretrizes.",
             ["14-novo-exercicio", "16-dieta"], [
    ("Sugestões de exercícios", "Seis exercícios do programa Vivifrail prontos para usar: um toque preenche o formulário.", GREEN_LIGHT, GREEN),
    ("Dias da semana", "Escolha os dias de cada exercício; o app recomenda de 3 a 5 vezes por semana.", PRIMARY_LIGHT, PRIMARY),
    ("Dieta da semana", "Café, almoço, lanche e jantar por dia, com botão para copiar o cardápio para a semana toda.", AMBER_LIGHT, AMBER),
    ("Lembrete de proteína", "Exemplos de alimentos proteicos e meta diária no topo da dieta.", AMBER_LIGHT, AMBER),
])

slide = screen_slide("Configurações", "Personalização e acessibilidade",
                     "Tudo ajustável pelo cuidador, sem internet.",
                     ["17-configuracoes", "18-configuracoes-saude", "20-alto-contraste"], [
    ("Perfis com nome e foto", "Foto escolhida da galeria, recortada e guardada no aparelho.", GREEN_LIGHT, GREEN),
    ("Saúde do idoso", "Peso, fragilidade e meta de água.", AMBER_LIGHT, AMBER),
    ("Contatos de SOS", "Número de emergência e telefone do cuidador.", RED_LIGHT, RED),
    ("Acessibilidade", "Fonte de 80% a 150%, alto contraste e animações que podem ser desligadas.", PRIMARY_LIGHT, PRIMARY),
])

# ==================== 13. MELHORIAS ====================
s = new_slide()
header(s, "Nesta versão", "Melhorias feitas a partir da pesquisa",
       "Cada recomendação clínica virou uma funcionalidade concreta no app.")
rows = [
    ("Hipodipsia: idosos sentem menos sede", "Contador de copos de água com meta diária e dica de hidratação"),
    ("Proteína: 1,0–1,5 g/kg/dia (ESPEN, PROT-AGE)", "Meta de proteína calculada pelo peso e fragilidade; refeições do dia com marcação"),
    ("Treino multicomponente 3–5x/semana (Vivifrail)", "Exercícios por tipo e dias da semana, sugestões prontas e meta semanal"),
    ("Risco de queda e tontura no exercício", "Pergunta pós-exercício, botão \"Tive uma queda\" e alertas para o cuidador"),
    ("Saúde afetiva (GDS-15)", "Check-in diário de humor e sintomas, com histórico da semana"),
    ("Adesão sobe com monitoramento digital", "Relatório de adesão dos últimos 7 dias no painel do cuidador"),
    ("Barreira cognitiva e visual (MOLD-US)", "Alto contraste, fonte ampliável, botões grandes e modais simples"),
    ("Segurança e LGPD", "Senha guardada como hash e todos os dados apenas no aparelho"),
]
y = 2.15
rect(s, 0.7, y, 11.9, 0.45, fill=PRIMARY, radius=0.08)
text(s, 0.95, y + 0.1, 5.2, 0.3, "Evidência da pesquisa", size=13, bold=True, color=WHITE)
text(s, 6.45, y + 0.1, 5.9, 0.3, "Melhoria no CuidarApp", size=13, bold=True, color=WHITE)
y += 0.5
for i, (evidence, improvement) in enumerate(rows):
    rect(s, 0.7, y, 11.9, 0.46, fill=WHITE if i % 2 == 0 else PRIMARY_LIGHT, radius=0.06)
    text(s, 0.95, y + 0.11, 5.3, 0.3, evidence, size=12, bold=True)
    text(s, 6.45, y + 0.11, 6.0, 0.3, improvement, size=12, color=TEXT_2)
    y += 0.5
text(s, 0.7, 6.65, 11.9, 0.3, "Também corrigido: edição de nomes, senha e telefones (o WebView não exibia as janelas antigas), botão voltar do Android, visitas do dia e navegação inferior.",
     size=11, color=TEXT_2)
footer(s)

# ==================== 14. BENEFÍCIOS ====================
s = new_slide()
header(s, "Benefícios", "Por que o CuidarApp faz diferença")
cols = [
    ("Para o idoso", GREEN, GREEN_LIGHT, [
        "Mais autonomia: sabe o que fazer a cada hora do dia",
        "Menos esquecimentos de remédios, água e refeições",
        "Segurança: SOS e registro de queda a um toque",
        "Tela simples, letras grandes e alto contraste",
    ]),
    ("Para o cuidador", PRIMARY, PRIMARY_LIGHT, [
        "Acompanha a rotina sem precisar perguntar tudo",
        "Alertas de queda, mal-estar e humor baixo",
        "Relatório semanal pronto para levar à consulta",
        "Cadastro rápido com sugestões baseadas em diretrizes",
    ]),
    ("Para a saúde", BLUE, BLUE_LIGHT, [
        "Nutrição + exercício guiado: força de preensão +24,8%",
        "Velocidade de marcha +19,6%",
        "Quedas até 34% menores com protocolo multicomponente",
        "Adesão de 78,4% com monitoramento digital",
    ]),
]
cw = 3.85
for i, (title, strong, light, items) in enumerate(cols):
    x = 0.7 + i * (cw + 0.17)
    rect(s, x, 1.85, cw, 4.75, fill=WHITE, line=BORDER)
    rect(s, x, 1.85, cw, 0.75, fill=light, radius=0.12)
    text(s, x + 0.3, 2.02, cw - 0.6, 0.45, title, size=20, bold=True, color=strong)
    cy = 2.9
    for item in items:
        circle(s, x + 0.32, cy + 0.08, 0.16, strong)
        text(s, x + 0.62, cy, cw - 0.9, 0.8, item, size=14)
        cy += 0.9
text(s, 0.7, 6.7, 11.9, 0.3, "Indicadores clínicos: meta-análise com 1.420 idosos após 24 semanas (Tieland et al., 2012; Sherrington et al., 2019; Izquierdo et al., 2019).",
     size=10, color=TEXT_2)
footer(s)

# ==================== 15. COMPARATIVO ====================
s = new_slide()
header(s, "Diferenciais", "Comparativo com aplicativos semelhantes")
table_rows = [
    ("", "Nutrição e água", "Exercício", "Remédios", "Offline", "Foco no idoso"),
    ("Vivifrail", "—", "●", "—", "◐", "●"),
    ("MyFitnessPal", "●", "◐", "—", "—", "—"),
    ("Medisafe / CareZone", "—", "—", "●", "—", "◐"),
    ("Meu SUS Digital", "—", "—", "◐", "—", "◐"),
    ("CuidarApp", "●", "●", "●", "●", "●"),
]
col_w = [3.1, 1.76, 1.76, 1.76, 1.76, 1.76]
y = 1.95
for r, row in enumerate(table_rows):
    is_head = r == 0
    is_us = row[0] == "CuidarApp"
    fill = PRIMARY if is_us else (WHITE if r % 2 else PRIMARY_LIGHT) if not is_head else BG
    if not is_head:
        rect(s, 0.7, y, sum(col_w), 0.62, fill=fill, radius=0.08)
    x = 0.7
    for c, cell in enumerate(row):
        color = WHITE if is_us else (TEXT_2 if is_head else (TEXT if c == 0 else PRIMARY))
        text(s, x + (0.25 if c == 0 else 0), y + 0.14, col_w[c] - (0.25 if c == 0 else 0), 0.35, cell,
             size=13 if is_head else (15 if c == 0 else 18), bold=is_head or c == 0 or is_us, color=color,
             align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
        x += col_w[c]
    y += 0.7
text(s, 0.7, 6.2, 11.9, 0.5, "● completo   ◐ parcial   — ausente.  Fonte: benchmarking da pesquisa (Wildenbos et al.; Izquierdo et al.).",
     size=11, color=TEXT_2)
footer(s)

# ==================== 16. PRIVACIDADE ====================
s = new_slide()
header(s, "Privacidade", "Dados de saúde protegidos por padrão",
       "Registros de alimentação, hidratação e mobilidade são dados sensíveis (LGPD, Lei nº 13.709/2018, Art. 11).")
items = [
    ("Nada vai para a internet", "O app não pede permissão de internet. Tudo é gravado no armazenamento interno do aparelho.", GREEN_LIGHT, GREEN),
    ("Sem cadastro, sem conta", "Não há e-mail, CPF ou login online. Só nomes e, se quiser, uma foto.", BLUE_LIGHT, BLUE),
    ("Área do cuidador protegida", "Senha de 4 dígitos guardada como hash (SHA-256), nunca em texto puro.", PRIMARY_LIGHT, PRIMARY),
    ("Controle total", "O cuidador pode apagar todos os dados a qualquer momento nas configurações.", RED_LIGHT, RED),
]
for i, (title, body, light, strong) in enumerate(items):
    col, row = i % 2, i // 2
    x, y = 0.7 + col * 6.05, 2.45 + row * 2.05
    rect(s, x, y, 5.85, 1.8, fill=WHITE, line=BORDER)
    circle(s, x + 0.3, y + 0.35, 0.55, light)
    text(s, x + 0.3, y + 0.35, 0.55, 0.55, "✓", size=20, bold=True, color=strong, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.1, y + 0.33, 4.5, 0.4, title, size=17, bold=True)
    text(s, x + 1.1, y + 0.8, 4.5, 0.9, body, size=13, color=TEXT_2)
footer(s)

# ==================== 17. ENCERRAMENTO ====================
s = new_slide(bg=PRIMARY)
circle(s, -1.5, 4.3, 5.0, PRIMARY_DARK)
circle(s, 10.5, -1.8, 5.0, RGBColor(0x8F, 0x74, 0xFD))
text(s, 0.9, 1.3, 7.5, 0.9, "Como começar", size=40, bold=True, color=WHITE)
steps = [
    "1. Instale o APK no celular Android (versão 8.0 ou superior).",
    "2. Toque em \"Cuidador\" e entre com a senha 1234.",
    "3. Em Configurações, defina nomes, fotos, peso, telefones e troque a senha.",
    "4. Cadastre remédios, exercícios sugeridos, dieta e visitas.",
    "5. Entregue o celular ao idoso: ele entra com um toque.",
]
text(s, 0.9, 2.4, 7.0, 3.2, [(st, {"space_after": 10}) for st in steps], size=17, color=WHITE)
text(s, 0.9, 5.9, 7.0, 0.6, "CuidarApp — cuidar bem, todos os dias.", size=22, bold=True, color=WHITE)
phone(s, "11-cuidador-painel", 9.0, 0.9, 5.8)

prs.save(OUTPUT)
print("Apresentação salva em", OUTPUT, "-", SLIDE_COUNT[0], "slides")
